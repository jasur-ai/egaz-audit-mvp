#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""TZ-1 T1 paketini Word hujjat sifatida yasash (chop etish va imzolash uchun).

Manba: `YAKUNIY/9-TZ-1-T1-PAKETI.md` · `YAKUNIY/6-TZ-1-Shovqin-Pilot.md` §4–§6.
Chiqish: `YAKUNIY/12-T1-XAT-VA-MEMORANDUM.docx` — 4 qism:
  1) Xat (korxona rahbariga)   2) Memorandum (imzo uchun)
  3) Ma'lumot ilovasi (D1–D4)  4) Obyekt kartasi (M1–M5, B1–B6) + kalendar

Ishlatish:
  python3 scripts/build_t1_packet.py                       # standart yo'lga yozadi
  python3 scripts/build_t1_packet.py --out /tmp/t1.docx
  python3 scripts/build_t1_packet.py --company "«XXX» IES" --date 2026-10-12
"""
from __future__ import annotations

import argparse
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

OUT_DEFAULT = "/home/user/YAKUNIY/12-T1-XAT-VA-MEMORANDUM.docx"

# --- to'ldiriladigan maydonlar (korxona nomi, sanalar)
FIELDS = {
    "kompaniya": "⟦Korxona nomi⟧",
    "obyekt": "⟦Obyekt (IES/IEM/zavod)⟧",
    "rahbar": "⟦Rahbar F.I.Sh., lavozim⟧",
    "muallif": "⟦Tadqiqot guruhi rahbari F.I.Sh.⟧",
    "aloqa": "⟦telefon, e-pochta⟧",
}


def _h(doc, text, level=1):
    doc.add_heading(text, level=level)


def _p(doc, text, bold=False, italic=False, size=None):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = size
    return par


def _table(doc, rows, header=True, widths=None):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Light Grid Accent 1" if header else "Table Grid"
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for c, val in zip(cells, row):
            c.text = str(val)
            if i == 0 and header:
                for par in c.paragraphs:
                    for run in par.runs:
                        run.bold = True
    if widths:
        for r in t.rows:
            for c, w in zip(r.cells, widths):
                c.width = w
    return t


# ------------------------------------------------------------------ 1) xat

def build_letter(doc, f: dict, sana: str) -> None:
    _h(doc, "1. XAT", 1)
    _p(doc, f"{f['kompaniya']} rahbari {f['rahbar']}ga", bold=True)
    _p(doc, f"Toshkent sh. · {sana} · № ⟦chiqish raqami⟧", italic=True)
    doc.add_paragraph()
    _p(doc, "HURMATLI ⟦Rahbar ismi⟧!", bold=True)
    _p(doc, "O'lchov ishonchliligi bo'yicha tadqiqot guruhi «Shovqin qavati» pilotini o'tkazadi. "
            "Pilotning maqsadi — korxonaning hisobot qiymati bilan avtomatik monitoring stansiyasi (CEMS) "
            "qiymati orasidagi farq taqsimotini o'lchash va shu asosda o'lchov ishonchliligini baholash "
            "mezonlarini (qabul / shartli / rad) asoslash.")
    _p(doc, f"Tadqiqot {f['obyekt']} misolida o'tkazilishi rejalashtirilgan. Ishtirok quyidagilarni nazarda tutadi:")
    for x in ("D1: CEMS 20 daqiqalik o'rtacha qiymatlari (kunlik fayl);",
              "D2: hisobot qiymati (xuddi shu oyna);",
              "D3: uskuna holati jurnali (to'xtash / ta'mir / kalibrovka);",
              "D4: oqim o'lchagichi hujjatlari (turi, joylashuvi, oxirgi tekshiruv natijasi)."):
        doc.add_paragraph(x, style="List Bullet")
    _p(doc, "Korxona uchun kafolatlar: pilot — tekshiruv yoki jarima emas; natija korxonaga nashrdan 5 ish kuni "
            "oldin beriladi; raqamlar kod ostida (K-1, K-2, K-3) e'lon qilinadi; xom qatorlar ommaga chiqmaydi; "
            "shaxs ma'lumotlari yig'ilmaydi; savdo siri so'ralmaydi.")
    _p(doc, "Iltimos, ilova qilingan memorandum loyihasini ko'rib chiqib, mas'ul mutaxassisni belgilang. "
            "Aloqa: " + f["aloqa"] + ".")
    doc.add_paragraph()
    _p(doc, f"{f['muallif']}  ______________", italic=True)


# ------------------------------------------------------------------ 2) memorandum

def build_memorandum(doc, f: dict, sana: str) -> None:
    doc.add_page_break()
    _h(doc, "2. MEMORANDUM", 1)
    _p(doc, "Shovqin qavatini o'lchash pilotida ishtirok etish to'g'risida", bold=True)
    _p(doc, f"⟦Shahar⟧, {sana}", italic=True)
    doc.add_paragraph()
    _p(doc, "Tomonlar:", bold=True)
    doc.add_paragraph(f"1. Tadqiqot guruhi — {f['muallif']} (keyingi o'rinlarda «Tadqiqot guruhi»), "
                      "ilmiy-tadqiqot maqsadida pilotni o'tkazuvchi tomon.", style="List Number")
    doc.add_paragraph(f"2. {f['kompaniya']} (keyingi o'rinlarda «Korxona»), {f['obyekt']} ob'ektining "
                      "egasi/mas'ul operatori.", style="List Number")
    doc.add_paragraph()

    items = [
        ("1. Memorandum predmeti",
         "Tomonlar hisobot qiymati bilan CEMS qiymati orasidagi farq taqsimotini o'lchash va shu asosda "
         "o'lchov ishonchliligini baholash mezonlarini (qabul / shartli / rad) asoslash maqsadida hamkorlik "
         "qiladi. Pilot natijasi — metodik tavsiyalar, tekshiruv yoki jarima emas."),
        ("2. Korxona nima beradi",
         "D1: CEMS 20 daqiqalik o'rtacha (kunlik fayl) · D2: hisobot qiymati (xuddi shu oyna) · "
         "D3: uskuna holati jurnali · D4: oqim o'lchagichi hujjatlari. Batafsil — Ma'lumot ilovasi (3-qism)."),
        ("3. Korxona nima oladi",
         "Pilot natijasining to'liq nusxasi nashrdan 5 ish kuni oldin · raqamlar kod ostida e'lon qilinadi "
         "(nom faqat yozma rozilik bilan) · hisobotning metodik qismi korxona uchun tekin."),
        ("4. Nima talab qilinmaydi",
         "Shaxs ma'lumotlari yig'ilmaydi · xom qatorlar ommaga chiqmaydi (faqat taqsimot, kvantil, MAD) · "
         "detektor/filtr samarasi baholanmaydi · savdo siri so'ralmaydi."),
        ("5. Muddat va tartib",
         "Ma'lumot yig'ish oynasi 07.12.2026 – 29.01.2027 (8 hafta, zaxira +4 hafta); bosqichlar TZ-1 §6 "
         "(T1–T7) bo'yicha; aloqa — tayinlangan mas'ullar orqali."),
        ("6. Maxfiylik",
         "Ma'lumot ikki nusxada (ishchi + zaxira), append-only; o'zgartirish taqiqlanadi, tuzatish — yangi "
         "versiya sifatida; pilot yakunida korxona talabiga ko'ra ma'lumot qaytariladi/o'chiriladi (arxiv muddati — kelishuvga ko'ra)."),
        ("7. Bekor qilish",
         "Har bir tomon 15 kun oldin yozma xabar bilan pilotdan chiqishi mumkin; shu paytgacha olingan "
         "ma'lumot korxona talabiga ko'ra qaytariladi."),
        ("8. Yakuniy band",
         "Memorandum majburiy shartnoma emas — niyat va ma'lumot almashish tartibini belgilaydi; moliyaviy "
         "majburiyat tugdirmaydi."),
    ]
    for head, body in items:
        _p(doc, head, bold=True)
        _p(doc, body)
        doc.add_paragraph()

    doc.add_paragraph()
    _table(doc, [["Tadqiqot guruhi", "Korxona"],
                 [f"{f['muallif']}\n\n______________", f"{f['rahbar']}\n\n______________"]], header=False)
    _p(doc, "Ilova: Ma'lumot ilovasi (3-qism) va Obyekt kartasi (4-qism, ichki).", italic=True)


# ------------------------------------------------------------------ 3) ma'lumot ilovasi

def build_annex(doc) -> None:
    doc.add_page_break()
    _h(doc, "3. MA'LUMOT ILOVASI", 1)
    _p(doc, "Memorandumga qo'shiladi. Qabul shartlari TZ-1 §4.2/§4.3 bilan bir xil.", italic=True)
    doc.add_paragraph()

    _h(doc, "3.1. Oqimlar (D1–D4)", 2)
    _table(doc, [
        ["Kod", "Ma'lumot", "Kim beradi", "Kadans", "Birinchi yuborish"],
        ["D1", "CEMS 20 daqiqalik o'rtacha", "Korxona stansiyasi", "Kunlik fayl", "⟦sana⟧"],
        ["D2", "Hisobot qiymati (xuddi shu oyna)", "Ekologiya xizmati", "Kunlik fayl", "⟦sana⟧"],
        ["D3", "Uskuna jurnali (to'xtash/ta'mir/kalibrovka)", "Texnik xizmat", "Haftalik", "⟦sana⟧"],
        ["D4", "Oqim o'lchagichi: turi, joylashuvi, tekshiruv", "Metrologiya xizmati", "Bir marta", "⟦sana⟧"],
    ])
    doc.add_paragraph()

    _h(doc, "3.2. Fayl sxemasi (qabul sharti)", 2)
    _p(doc, "Kalit: (obyekt_id, modda, oyna_boshi, manba) — takrorlanish bo'lsa fayl RAD ETILADI. "
            "Vaqt: ISO-8601, mahalliy UTC+5 offseti bilan (masalan 2026-12-07T08:00:00+05:00). "
            "Oyna boshi 20 daqiqaga karrali bo'lishi shart.")
    _table(doc, [
        ["Kolonka", "Tur", "Majburiy", "Izoh"],
        ["obyekt_id", "matn", "✅", "Shartnomada berilgan kod (nomi emas)"],
        ["modda", "matn", "✅", "SO2 / NOx / CO / PM"],
        ["oyna_boshi", "vaqt", "✅", "20 daqiqalik oyna boshlanishi (offset bilan)"],
        ["qiymat", "son", "✅", "Bo'sh bo'lsa — valid=false bilan yoziladi"],
        ["birlik", "matn", "✅", "mg/m³ (SO₂, NOx, CO) · µg/m³ (PM) · m³/s (oqim)"],
        ["valid", "bool", "✅", "false = uskuna to'xtagan/norasmiy oyna"],
        ["manba", "matn", "✅", "cems | hisobot | etalon"],
        ["izoh", "matn", "➖", "Kalibrovka/to'xtash belgisi (D3 bilan bog'lanadi)"],
    ])
    doc.add_paragraph()

    _h(doc, "3.3. Sifat nazorati (avtomatik tekshiriladi)", 2)
    for x in ("Yetishmayotgan oyna > 10% (fayl qamrovi ichida) → obyekt «shartli»ga o'tadi, sabab D3 dan izohlanadi;",
              "Takroriy kalit → fayl qaytariladi; birlik moddaga mos emas → qaytariladi;",
              "Har fayl uchun audit izi: SHA-256 · yuklash vaqti · kim yukladi (rol) · qatorlar soni;",
              "Ma'lumot append-only saqlanadi: o'zgartirish taqiqlanadi, tuzatish — yangi versiya sifatida."):
        doc.add_paragraph(x, style="List Bullet")


# ------------------------------------------------------------------ 4) obyekt kartasi

def build_object_card(doc) -> None:
    doc.add_page_break()
    _h(doc, "4. OBYEKT KARTASI (ichki hujjat — korxonaga yuborilmaydi)", 1)

    _h(doc, "4.1. Majburiy mezonlar (biri bajarilmasa — obyekt olinmaydi)", 2)
    _table(doc, [
        ["#", "Mezon", "Tekshirish usuli", "Natija"],
        ["M1", "Yozma rozilik (memorandum)", "Imzolangan hujjat", "⟦sana⟧"],
        ["M2", "CEMS mavjud va ishlaydi", "O'rnatish dalili + 90 kun uzluksizlik", "⟦%⟧"],
        ["M3", "Hisobot O'z DSt 3605:2022 bo'yicha", "Hisobot shakli", "⟦ha/yo'q⟧"],
        ["M4", "Texnik jurnal (6 oy)", "Nusxa", "⟦ha/yo'q⟧"],
        ["M5", "Aloqa kanali (2 mutaxassis)", "Yozma ro'yxat", "⟦F.I.Sh.⟧"],
    ])
    doc.add_paragraph()

    _h(doc, "4.2. Baholanadigan mezonlar (0–3 ball · tavsiya: jami ≥12/18)", 2)
    _table(doc, [
        ["#", "Mezon", "0 ball", "3 ball", "Ball"],
        ["B1", "Hisobot ↔ fiskal (gaz/elektr) solishtirish", "yo'q", "agregat darajada bor", "⟦ ⟧"],
        ["B2", "Laboratoriya bilan parallel o'lchov tajribasi", "yo'q", "doimiy amaliyot", "⟦ ⟧"],
        ["B3", "Yoqilg'i tarkibi o'zgaruvchanligi hujjatlashgan", "yo'q", "oylik tahlil", "⟦ ⟧"],
        ["B4", "Uskuna/yondashuv almashinuvi tarixi", "yo'q", "so'nggi 3 yilda", "⟦ ⟧"],
        ["B5", "Oqim o'lchagichi hujjatlari to'liq", "yo'q", "X-shakl/RATA izi bor", "⟦ ⟧"],
        ["B6", "Oldingi inspeksiya/audit natijasi", "yo'q", "rasmiy hisobot", "⟦ ⟧"],
    ])
    doc.add_paragraph()
    _p(doc, "Xulosa: ⟦obyekt⟧ → S1 / S2 / S3 stratasiga: ☐ tanlandi  ☐ zaxira  ☐ rad", bold=True)
    doc.add_paragraph()

    _h(doc, "4.3. T1 kalendari (12.10 – 06.11.2026)", 2)
    _table(doc, [
        ["Sana", "Qadam", "Mas'ul", "Natija"],
        ["12.10.2026", "Rasmiy xat + memorandum loyihasi yuboriladi", "Muallif", "Yuborish qayd etiladi"],
        ["13.10–18.10", "Mas'ul mutaxassis aniqlanadi (M5)", "Muallif", "2 ism + aloqa kanali"],
        ["19.10–25.10", "Obyekt kartasi to'ldiriladi (M1–M5, B1–B6)", "Tadqiqot guruhi", "Nomzodlar baholandi"],
        ["26.10–02.11", "Uchrashuv: memorandum muhokamasi va imzo", "Muallif + korxona", "S1 imzo"],
        ["03.11–06.11", "Qolgan imzolar + T2/T3 ga start", "Muallif + brigada", "3 imzo"],
    ])
    doc.add_paragraph()
    _p(doc, "Eslatma: xat va memorandum imzolangach, ma'lumot fayllari qabul moduli bilan tekshiriladi "
            "(scripts/validate_pilot_data.py) — xatolar bo'lsa fayl qaytariladi va sabab yoziladi.", italic=True)


def main() -> int:
    ap = argparse.ArgumentParser(description="T1 paketi (xat + memorandum) DOCX generatori")
    ap.add_argument("--out", default=OUT_DEFAULT)
    ap.add_argument("--company", default=FIELDS["kompaniya"])
    ap.add_argument("--obyekt", default=FIELDS["obyekt"])
    ap.add_argument("--date", default=date.today().isoformat())
    args = ap.parse_args()

    try:
        import docx  # noqa: F401
    except Exception:
        print("❌ python-docx o'rnatilmagan: pip install python-docx")
        return 2

    from docx import Document

    f = dict(FIELDS)
    f["kompaniya"], f["obyekt"] = args.company, args.obyekt

    doc = Document()
    doc.core_properties.title = "TZ-1 T1 paketi — xat va memorandum"
    doc.core_properties.author = f["muallif"]

    _h(doc, "TZ-1 — T1 PAKETI (xat, memorandum, ma'lumot ilovasi, obyekt kartasi)", 0)
    _p(doc, f"Shovqin qavatini o'lchash piloti · {args.date} · TZ-1 v1.0 asosida "
            f"(YAKUNIY/6-TZ-1-Shovqin-Pilot.md)", italic=True)
    doc.add_paragraph()
    _p(doc, "Bu hujjat: (1) korxona rahbariga xat, (2) imzolanadigan memorandum, (3) texnik ilova — ma'lumot "
            "oqimlari va sxema, (4) ichki obyekt kartasi va kalendar. To'ldiriladigan joylar ⟦ ⟧ bilan belgilangan.")

    build_letter(doc, f, args.date)
    build_memorandum(doc, f, args.date)
    build_annex(doc)
    build_object_card(doc)

    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    doc.save(args.out)
    size = os.path.getsize(args.out)
    print(f"✅ T1 paketi: {args.out} ({size:,} bayt · bo'limlar: xat, memorandum, ilova, karta)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
